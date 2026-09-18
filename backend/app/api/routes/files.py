import base64
import io
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from typing import Any
from urllib.parse import quote

import httpx
import xlrd
from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import HTMLResponse, PlainTextResponse, StreamingResponse
from openpyxl import load_workbook
from pypdf import PdfReader

from app.ai.prompts import RESUME_PARSE_PROMPT
from app.ai.provider import AIClient
from app.ai.workflows import run_structured_completion
from app.api.dependencies import Session, WorkspaceOwner
from app.services.resume_export import render_resume
from app.services.resumes import create_resume, owned_resume, resume_dto

router = APIRouter(prefix="/api", tags=["files"])


def plain_text(resume: dict[str, Any]) -> str:
    lines: list[str] = []
    for section in resume.get("sections", []):
        if not section.get("visible", True):
            continue
        content = section.get("content") or {}
        if section.get("type") == "personal_info":
            lines.extend(
                str(content.get(key, "")) for key in ("fullName", "jobTitle") if content.get(key)
            )
            lines.append(
                " | ".join(
                    str(content.get(key, ""))
                    for key in ("email", "phone", "location")
                    if content.get(key)
                )
            )
        elif section.get("type") == "summary":
            lines.extend([f"== {section.get('title', '')} ==", str(content.get("text", ""))])
        elif section.get("type") == "skills":
            lines.append(f"== {section.get('title', '')} ==")
            for category in content.get("categories", []):
                lines.append(
                    f"- {category.get('name', '')}: {', '.join(category.get('skills', []))}"
                )
        else:
            lines.append(f"== {section.get('title', '')} ==")
            for item in content.get("items", []):
                title = (
                    item.get("position")
                    or item.get("name")
                    or item.get("degree")
                    or item.get("title")
                    or ""
                )
                subtitle = (
                    item.get("company") or item.get("institution") or item.get("subtitle") or ""
                )
                lines.append(f"- {title}{' - ' + subtitle if subtitle else ''}")
                if item.get("description"):
                    lines.append(f"  {item['description']}")
                for value in item.get("highlights", []):
                    lines.append(f"  * {value}")
        lines.append("")
    return "\n".join(lines)


def attachment_header(filename: str) -> str:
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "-", filename).strip("-") or "resume"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


@router.get("/resume/{resume_id}/export")
async def export_resume(
    resume_id: str,
    session: Session,
    user: WorkspaceOwner,
    format: str = "json",
    forPrint: bool = False,
    fitOnePage: bool = False,
):
    value = resume_dto(await owned_resume(session, user, resume_id))
    filename = f"{value['title']}-{datetime.now():%Y%m%d%H%M%S}"
    if format == "json":
        return value
    if format == "html":
        return HTMLResponse(
            await render_resume(value, "html", for_print=forPrint),
            headers={"Content-Disposition": attachment_header(f"{filename}.html")},
        )
    if format == "txt":
        return PlainTextResponse(
            plain_text(value),
            headers={"Content-Disposition": attachment_header(f"{filename}.txt")},
        )
    if format in {"docx", "pdf"}:
        output = await render_resume(value, format, fit_one_page=fitOnePage)
        return StreamingResponse(
            io.BytesIO(output),
            media_type="application/pdf" if format == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": attachment_header(f"{filename}.{format}")},
        )
    raise HTTPException(status_code=400, detail="Unsupported export format")


async def file_text(file: UploadFile, content: bytes) -> str:
    content_type = file.content_type or "application/octet-stream"
    name = file.filename or "file"
    lower_name = name.lower()
    if content_type.startswith("text/") or lower_name.endswith(
        (".txt", ".md", ".markdown", ".csv", ".json", ".xml", ".yaml", ".yml", ".log", ".rtf")
    ):
        if content.startswith(b"\xff\xfe"):
            value = content[2:].decode("utf-16le", errors="replace")
        elif content.startswith(b"\xfe\xff"):
            value = content[2:].decode("utf-16be", errors="replace")
        else:
            value = content.decode("utf-8-sig", errors="replace")
        return value.replace("\x00", "")[:120_000]
    if content_type == "application/pdf" or name.lower().endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if lower_name.endswith(".docx"):
        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)[:120_000]
    if lower_name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for worksheet in workbook.worksheets:
            lines.append(f"## 工作表：{worksheet.title}")
            for row in worksheet.iter_rows(values_only=True):
                cells = [
                    str(value or "").replace("|", "\\|").replace("\n", "<br>") for value in row
                ]
                if any(cells):
                    lines.append(f"| {' | '.join(cells)} |")
                if sum(len(value) for value in lines) >= 120_000:
                    return "\n".join(lines)[:119_980] + "\n[文件内容过长，已截断]"
        return "\n".join(lines)
    if lower_name.endswith(".xls"):
        workbook = xlrd.open_workbook(file_contents=content)
        lines = []
        for worksheet in workbook.sheets():
            lines.append(f"## 工作表：{worksheet.name}")
            for row_index in range(worksheet.nrows):
                cells = [
                    str(worksheet.cell_value(row_index, column))
                    .replace("|", "\\|")
                    .replace("\n", "<br>")
                    for column in range(worksheet.ncols)
                ]
                if any(cells):
                    lines.append(f"| {' | '.join(cells)} |")
        return "\n".join(lines)[:120_000]
    if lower_name.endswith(".doc"):
        antiword = shutil.which("antiword")
        if not antiword:
            raise HTTPException(
                status_code=400, detail="解析 .doc 文件需要安装 antiword；建议另存为 .docx"
            )
        with tempfile.NamedTemporaryFile(suffix=".doc") as source:
            source.write(content)
            source.flush()
            result = subprocess.run(
                [antiword, source.name], capture_output=True, check=False, timeout=30
            )
        if result.returncode != 0:
            raise HTTPException(status_code=400, detail="无法解析该 .doc 文件")
        return result.stdout.decode("utf-8", errors="replace")[:120_000]
    raise HTTPException(status_code=400, detail="不支持的文件类型")


@router.post("/ai/attachments/parse")
async def parse_attachment(user: WorkspaceOwner, file: UploadFile = File(...)) -> dict:
    del user
    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="文件不能超过 20MB")
    return {
        "content": await file_text(file, content),
        "contentType": file.content_type or "application/octet-stream",
        "name": file.filename or "file",
        "size": len(content),
    }


@router.post("/resume/parse", status_code=201)
async def parse_resume(
    session: Session,
    user: WorkspaceOwner,
    file: UploadFile = File(...),
    template: str = Form("classic"),
    language: str = Form("zh"),
) -> dict:
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large")
    try:
        text = await file_text(file, content)
    except HTTPException:
        if not (file.content_type or "").startswith("image/"):
            raise
        parsed = await run_structured_completion(
            AIClient(),
            system=RESUME_PARSE_PROMPT,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": "解析这张简历"},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:{file.content_type};base64,{base64.b64encode(content).decode()}"
                            },
                        },
                    ],
                },
            ],
            max_tokens=16384,
        )
    else:
        parsed = await run_structured_completion(
            AIClient(),
            system=RESUME_PARSE_PROMPT,
            prompt=text,
            max_tokens=16384,
        )
    mapping = [
        ("personal_info", "个人信息", parsed.get("personalInfo", {})),
        ("summary", "个人简介", {"text": parsed.get("summary", "")}),
        ("work_experience", "工作经历", {"items": parsed.get("workExperience", [])}),
        ("education", "教育背景", {"items": parsed.get("education", [])}),
        ("skills", "专业技能", {"categories": parsed.get("skills", [])}),
        ("projects", "项目经历", {"items": parsed.get("projects", [])}),
        ("certifications", "证书", {"items": parsed.get("certifications", [])}),
        ("languages", "语言", {"items": parsed.get("languages", [])}),
    ]
    resume = await create_resume(
        session,
        user,
        {
            "title": parsed.get("personalInfo", {}).get("fullName") or "未命名简历",
            "template": template,
            "language": language,
            "sections": [
                {"type": kind, "title": title, "content": value} for kind, title, value in mapping
            ],
        },
    )
    return resume_dto(resume)


@router.get("/github/repo")
async def github_repo(url: str) -> dict:
    match = re.fullmatch(
        r"https?://github\.com/([A-Za-z0-9_.-]+)/([A-Za-z0-9_.-]+?)(?:\.git)?/?", url
    )
    if not match:
        raise HTTPException(status_code=400, detail="Invalid GitHub URL")
    owner, repo = match.groups()
    async with httpx.AsyncClient(timeout=20, trust_env=False) as client:
        response = await client.get(
            f"https://api.github.com/repos/{owner}/{repo}",
            headers={"Accept": "application/vnd.github+json"},
        )
    if response.status_code == 404:
        raise HTTPException(status_code=404, detail="Repository not found")
    if response.status_code == 403:
        raise HTTPException(status_code=429, detail="GitHub API rate limit exceeded")
    response.raise_for_status()
    data = response.json()
    return {
        "name": data["full_name"],
        "stars": data["stargazers_count"],
        "language": data.get("language") or "",
        "description": data.get("description") or "",
        "url": data["html_url"],
    }


