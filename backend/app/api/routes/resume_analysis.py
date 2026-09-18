import io
import logging
from datetime import UTC, datetime
from urllib.parse import quote
from uuid import uuid4

from docx import Document
from fastapi import APIRouter, File, Form, HTTPException, Response, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError

from app.ai.provider import AIConfigurationError
from app.api.dependencies import Session, WorkspaceOwner
from app.api.routes.files import plain_text
from app.db.models import ResumeAnalysisReport
from app.db.session import SessionFactory
from app.services.resume_analysis import analyze_text
from app.services.resume_export import render_resume
from app.services.resumes import owned_resume, resume_dto

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/resume-analysis", tags=["resume-analysis"])


def extract_document(content: bytes, filename: str) -> tuple[str, int | None]:
    try:
        if filename.lower().endswith(".pdf"):
            pdf = PdfReader(io.BytesIO(content))
            if len(pdf.pages) > 30:
                raise HTTPException(422, "请上传不超过 30 页的简历")
            return "\n".join(page.extract_text() or "" for page in pdf.pages), len(pdf.pages)
        if filename.lower().endswith(".docx"):
            doc = Document(io.BytesIO(content))
            lines = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            for section in doc.sections:
                lines.extend(p.text for p in section.header.paragraphs)
                lines.extend(p.text for p in section.footer.paragraphs)
            return "\n".join(lines), None
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, "文件无法读取，请确认文件完整且未加密") from error
    raise HTTPException(422, "请上传 PDF 或 DOCX 简历")


@router.post("")
async def create_report(
    session: Session, user: WorkspaceOwner,
    file: UploadFile | None = File(None), resumeId: str = Form(""),
    role: str = Form("", max_length=200), jobDescription: str = Form("", max_length=15000),
) -> dict:
    owner_id = user.id
    if bool(file) == bool(resumeId):
        raise HTTPException(422, "请选择一个简历来源")
    sections = []
    pages = None
    if resumeId:
        resume = await owned_resume(session, user, resumeId)
        dto = resume_dto(resume)
        text = plain_text(dto)
        title = resume.title
        sections = [s["title"] for s in dto["sections"] if s.get("visible", True)]
    else:
        assert file is not None
        content = await file.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(413, "文件大小不能超过 10 MB")
        title = (file.filename or "上传简历")[:200]
        text, pages = extract_document(content, title)
    if len(text.strip()) < 40:
        raise HTTPException(422, "未识别到足够文字。扫描版 PDF 请先转换为可提取文字的文件。")
    if len(text) > 60000:
        raise HTTPException(422, "简历文字过长，请上传不超过 60000 字的文件")
    # Return the source transaction to the pool before waiting on the model.
    await session.rollback()
    try:
        analysis = await analyze_text(text, role.strip(), jobDescription.strip())
    except AIConfigurationError as error:
        raise HTTPException(503, str(error)) from error
    except ValueError as error:
        logger.error("Resume analysis has no complete report: category=%s", type(error).__name__)
        raise HTTPException(502, "本次未取得可用的分析结果，报告未生成。") from error
    except Exception as error:
        logger.error("Resume analysis failed: category=%s upstream_status=%s", type(error).__name__, getattr(error, "status_code", None))
        raise HTTPException(502, "AI 分析暂时不可用，请稍后重试。") from error
    identifier = str(uuid4())
    result = {
        "id": identifier, "title": f"{title} · 分析报告", "createdAt": datetime.now(UTC).isoformat(),
        "resumeId": resumeId or None, "target": role.strip() or analysis.get("targetRole", ""),
        "basics": {"charCount": len("".join(text.split())), "pageCount": pages, "sections": sections or analysis.get("sections", [])},
        "dimensions": analysis["dimensions"], "issues": analysis["issues"], "strengths": analysis["strengths"],
        "warnings": ["基于简历文字分析，未评价视觉版式；分数仅供完善文档参考。"] + analysis.get("warnings", []) + ([] if role.strip() or jobDescription.strip() or analysis.get("targetRole") else ["未提供目标岗位，岗位匹配暂不评分。"]),
        "schemaVersion": 1, "rubricVersion": "text-v1",
    }
    try:
        async with SessionFactory() as writer:
            writer.add(ResumeAnalysisReport(id=identifier, user_id=owner_id, result=result, source_text=text))
            await writer.commit()
    except SQLAlchemyError as error:
        logger.error("Resume report persistence failed: category=%s", type(error).__name__)
        raise HTTPException(503, "报告已生成，但保存服务暂时不可用，请稍后重试。") from error
    return result


@router.get("")
async def history(session: Session, user: WorkspaceOwner, response: Response) -> list[dict]:
    response.headers["Cache-Control"] = "private, no-store"
    rows = await session.scalars(select(ResumeAnalysisReport).where(
        ResumeAnalysisReport.user_id == user.id,
    ).order_by(ResumeAnalysisReport.created_at.desc()).limit(30))
    return [{**{k: row.result[k] for k in ("id", "title", "createdAt")}, "kind": row.result.get("kind", "resume")} for row in rows]


@router.get("/{report_id}")
async def get_report(report_id: str, session: Session, user: WorkspaceOwner, response: Response) -> dict:
    response.headers["Cache-Control"] = "private, no-store"
    row = await session.scalar(select(ResumeAnalysisReport).where(
        ResumeAnalysisReport.id == report_id, ResumeAnalysisReport.user_id == user.id,
    ))
    if row is None:
        raise HTTPException(404, "分析报告不存在")
    return row.result


@router.delete("/{report_id}")
async def delete_report(report_id: str, session: Session, user: WorkspaceOwner) -> dict:
    row = await session.scalar(select(ResumeAnalysisReport).where(
        ResumeAnalysisReport.id == report_id, ResumeAnalysisReport.user_id == user.id,
    ))
    if row is None:
        raise HTTPException(404, "分析报告不存在")
    await session.delete(row)
    await session.commit()
    return {"success": True}


class ReportPdfRequest(BaseModel):
    html: str = Field(min_length=1, max_length=8_000_000)


@router.post("/{report_id}/pdf")
async def download_report_pdf(report_id: str, body: ReportPdfRequest, session: Session, user: WorkspaceOwner):
    row = await session.scalar(select(ResumeAnalysisReport).where(
        ResumeAnalysisReport.id == report_id, ResumeAnalysisReport.user_id == user.id,
    ))
    if row is None and report_id != 'example':
        raise HTTPException(404, "分析报告不存在")
    title = str(row.result.get('title') or '分析报告')[:160] if row else '示例分析报告' 
    await session.rollback()
    pdf = await render_resume({'html': body.html}, 'report-pdf')
    if not pdf.startswith(b'%PDF-'):
        raise HTTPException(503, "PDF 生成失败，请稍后重试")
    return Response(pdf, media_type='application/pdf', headers={
        'Content-Disposition': f"attachment; filename*=UTF-8''{quote(title + '.pdf', safe='')}",
        'Cache-Control': 'private, no-store',
        'X-Content-Type-Options': 'nosniff',
    })
