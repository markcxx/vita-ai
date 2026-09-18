import io
from copy import deepcopy

import pytest
from docx import Document
from fastapi import HTTPException
from pydantic import ValidationError

from app.api.routes.resume_analysis import extract_document
from app.services.resume_analysis import LENGTHS, normalize_report


def fixture():
    return {"dimensions": {k: {"values": [80] * n, "summary": "有原文证据", "tags": []} for k, n in LENGTHS.items()}, "strengths": ["表达清楚"], "sections": ["工作经历"], "issues": [{"dimension": "trust", "title": "补充结果", "evidence": "负责产品设计", "suggestion": "补充实际交付成果", "rewrite": "", "priority": "high"}]}


def test_scores_missing_target_and_visual_evidence():
    raw = fixture()
    result = normalize_report(raw, "负责产品设计", False)
    assert result["dimensions"]["match"]["score"] is None
    assert result["dimensions"]["reading"]["values"][0] is None
    assert result["dimensions"]["get"]["score"] == 80
    assert raw["dimensions"]["match"]["values"] == [80] * 4


def test_reject_invalid_score_shape_and_evidence():
    for mutate in (lambda r: r["dimensions"]["trust"]["values"].__setitem__(0, 120), lambda r: r["dimensions"]["trust"]["values"].pop()):
        raw = deepcopy(fixture())
        mutate(raw)
        with pytest.raises((ValidationError, ValueError)):
            normalize_report(raw, "负责产品设计", True)


def test_docx_extracts_table_and_header():
    doc = Document()
    doc.add_paragraph("工作经历")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "负责产品设计"
    doc.sections[0].header.paragraphs[0].text = "求职简历"
    output = io.BytesIO()
    doc.save(output)
    text, pages = extract_document(output.getvalue(), "简历.docx")
    assert "负责产品设计" in text and "求职简历" in text
    assert pages is None


def test_bad_document_fails_instead_of_fake_report():
    with pytest.raises(HTTPException) as error:
        extract_document(b"not a pdf", "resume.pdf")
    assert error.value.status_code == 422


def test_extraction_whitespace_does_not_fail_report():
    raw = fixture()
    raw["issues"][0]["evidence"] = "负责产品设计"
    result = normalize_report(raw, "负责产品\n设计", False)
    assert result["issues"][0]["evidence"] == "负责产品\n设计"


def test_unverified_suggestion_does_not_discard_valid_report():
    raw = fixture()
    raw["issues"].append({**raw["issues"][0], "evidence": "伪造的原文"})
    raw["issues"].append({"title": "缺少格式的建议"})
    result = normalize_report(raw, "负责产品设计", False)
    assert len(result["issues"]) == 1
    assert result["discardedIssueCount"] == 2
    assert result["dimensions"]["get"]["score"] == 80


def test_numeric_strings_decimals_and_inferred_role():
    raw = fixture()
    raw["targetRole"] = "产品设计师"
    raw["dimensions"]["trust"]["values"] = ["80", 75.4, 80.0, None]
    result = normalize_report(raw, "求职意向：产品设计师；负责产品设计", False)
    assert result["targetRole"] == "产品设计师"
    assert result["dimensions"]["match"]["score"] == 80
    assert result["dimensions"]["trust"]["values"] == [80, 75, 80, None]
    raw["targetRole"] = "未提供的岗位"
    assert normalize_report(raw, "负责产品设计", False)["targetRole"] == ""


def test_structural_failure_is_repaired_automatically(monkeypatch):
    import asyncio
    from types import SimpleNamespace
    from app.services import resume_analysis
    calls = {}
    thinking = []
    class Client:
        settings = SimpleNamespace(ai_provider="openai", ai_thinking_mode="unsupported")
        async def structured_complete(self, **kwargs):
            key = next((k for k in resume_analysis.PART_INSTRUCTIONS if f"只生成 {k} 维度" in kwargs["system"]), "profile")
            calls[key] = calls.get(key, 0) + 1
            if key == "profile":
                return kwargs["schema"].model_validate({"targetRole": "", "strengths": [], "sections": ["工作经历"]})
            if key == "trust" and calls[key] == 1:
                raise ValueError("Structured response validation failed")
            return kwargs["schema"].model_validate({**fixture()["dimensions"][key], "issues": []})
    def make_client(*, thinking_enabled):
        thinking.append(thinking_enabled)
        return Client()
    monkeypatch.setattr(resume_analysis, "AIClient", make_client)
    result = asyncio.run(resume_analysis.analyze_text("负责产品设计"))
    assert calls["trust"] == 2
    assert calls["reading"] == 1
    assert calls["information"] == 1
    assert thinking == [False]
    assert len(result["dimensions"]) == 5
