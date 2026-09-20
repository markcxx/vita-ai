import io
from types import SimpleNamespace

import httpx
import pytest
from docx import Document
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook

from app.services import profile_import as service


@pytest.fixture
def anyio_backend():
    return 'asyncio'


def attachment(name, data):
    return UploadFile(filename=name, file=io.BytesIO(data))


def test_docx_tables_are_extracted():
    doc = Document()
    doc.add_paragraph('个人资料')
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = '姓名'
    row.cells[1].text = '测试用户'
    stream = io.BytesIO()
    doc.save(stream)
    assert '姓名 | 测试用户' in service.document_text(stream.getvalue(), '.docx')


def test_spreadsheet_preserves_zero_values():
    book = Workbook()
    book.active.append(['姓名', '测试用户', 0])
    stream = io.BytesIO()
    book.save(stream)
    assert '姓名 | 测试用户 | 0' in service.document_text(stream.getvalue(), '.xlsx')


def test_legacy_chinese_text_encoding():
    assert service.document_text('个人资料'.encode('gb18030'), '.txt') == '个人资料'


@pytest.mark.anyio
async def test_multiple_materials_produce_draft_with_server_ids(monkeypatch):
    monkeypatch.setattr(service, 'AIClient', lambda: SimpleNamespace(settings=SimpleNamespace(ai_provider='openai')))

    async def complete(client, **kwargs):
        content = kwargs['messages'][0]['content']
        assert '项目说明.txt' in content and '证书.md' in content
        return {'profile': {'personalInfo': {'fullName': '测试用户', 'github': 'https://github.com/demo'},
                            'certifications': [{'name': '证书'}]}, 'warnings': ['日期未明确']}

    monkeypatch.setattr(service, 'run_structured_completion', complete)
    result = await service.extract_profile([attachment('项目说明.txt', b'project material'), attachment('证书.md', b'certificate material')])
    assert result['profile']['isSample'] is False
    assert result['profile']['certifications'][0]['id']
    assert result['profile']['personalInfo']['links'][0]['url'] == 'https://github.com/demo'
    assert result['warnings'] == ['日期未明确']


@pytest.mark.anyio
@pytest.mark.parametrize(('name', 'content', 'status'), [('bad.exe', b'data', 422), ('empty.txt', b'', 413), ('bad.png', b'not an image', 422), ('big.txt', b'x' * (service.MAX_FILE + 1), 413)])
async def test_reject_invalid_upload_before_model(name, content, status):
    with pytest.raises(HTTPException) as error:
        await service.extract_profile([attachment(name, content)])
    assert error.value.status_code == status


@pytest.mark.anyio
async def test_unrelated_material_cannot_clear_profile(monkeypatch):
    monkeypatch.setattr(service, 'AIClient', lambda: SimpleNamespace(settings=SimpleNamespace(ai_provider='openai')))

    async def complete(*args, **kwargs):
        return {'profile': {}, 'warnings': ['不是个人资料']}

    monkeypatch.setattr(service, 'run_structured_completion', complete)
    with pytest.raises(HTTPException) as error:
        await service.extract_profile([attachment('notes.txt', b'unrelated material')])
    assert error.value.status_code == 422


@pytest.mark.anyio
async def test_route_requires_login():
    from app.config import get_settings
    from app.main import app

    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as client:
        response = await client.post('/api/profile/import', headers={'origin': get_settings().app_url.rstrip('/')}, files={'files': ('test.txt', b'material')})
    assert response.status_code == 401


@pytest.mark.anyio
async def test_stale_import_does_not_write_over_newer_profile(monkeypatch):
    from app.api.routes import resources
    from app.db.models import User

    record = SimpleNamespace(version=5, data={'summary': 'newer changes'})

    async def profile_for(*args):
        return record

    class Session:
        async def scalar(self, statement):
            return record

        async def commit(self):
            pytest.fail('A stale import must not commit')

    monkeypatch.setattr(resources, 'profile_for', profile_for)
    with pytest.raises(HTTPException) as error:
        await resources.save_profile(Session(), User(id='owner'), {'data': {'summary': 'old'}, 'expectedVersion': 4})
    assert error.value.status_code == 409
    assert record.data == {'summary': 'newer changes'}


@pytest.mark.anyio
async def test_image_is_sent_to_vision_model(monkeypatch):
    monkeypatch.setattr(service, 'AIClient', lambda: SimpleNamespace(settings=SimpleNamespace(ai_provider='openai')))

    async def complete(client, **kwargs):
        blocks = kwargs['messages'][0]['content']
        assert blocks[1]['image_url']['url'].startswith('data:image/png;base64,')
        return {'profile': {'personalInfo': {'fullName': '测试用户'}}}

    monkeypatch.setattr(service, 'run_structured_completion', complete)
    result = await service.extract_profile([attachment('certificate.png', b'\x89PNG\r\n\x1a\nmock')])
    assert result['profile']['personalInfo']['fullName'] == '测试用户'


@pytest.mark.anyio
async def test_disconnected_import_cancels_model_and_closes_files(monkeypatch):
    import asyncio

    from fastapi import Response

    from app.api.routes import profile_import as route

    cancelled = asyncio.Event()

    async def extract(files):
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.set()

    class Session:
        async def rollback(self):
            pass

    class Request:
        async def is_disconnected(self):
            return True

    monkeypatch.setattr(route, 'extract_profile', extract)
    file = attachment('notes.txt', b'material')
    with pytest.raises(HTTPException) as error:
        await route.import_profile(None, Session(), Request(), Response(), [file])
    assert error.value.status_code == 499
    assert cancelled.is_set()
    assert file.file.closed


@pytest.mark.anyio
async def test_empty_placeholder_entries_are_not_importable(monkeypatch):
    monkeypatch.setattr(service, 'AIClient', lambda: SimpleNamespace(settings=SimpleNamespace(ai_provider='openai')))

    async def complete(*args, **kwargs):
        return {'profile': {'experiences': [{}], 'skills': [{'skills': [' ']}], 'summary': ' '}}

    monkeypatch.setattr(service, 'run_structured_completion', complete)
    with pytest.raises(HTTPException) as error:
        await service.extract_profile([attachment('notes.txt', b'unrelated material')])
    assert error.value.status_code == 422
